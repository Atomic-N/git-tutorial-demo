import json
import subprocess
import sys
import time
from pathlib import Path

from openpyxl import load_workbook


EXCEL_PATH = Path(r"D:\other\AI\re\linux_do_interesting_posts_readable.xlsx")
SHEET_NAME = "精选视图"
CACHE_DIR = Path(r"D:\other\AI\project\AI_TEF\data\linuxdo-json-cache")
AGENT_BROWSER_CLI = r"C:\Users\yzt82\AppData\Roaming\npm\agent-browser-cli.cmd"


def run_cmd(args):
    completed = subprocess.run(args, capture_output=True, text=True, encoding="utf-8", timeout=90)
    if completed.returncode != 0:
        raise RuntimeError(completed.stderr.strip() or completed.stdout.strip() or f"Command failed: {args}")
    return completed.stdout


def read_links(limit):
    wb = load_workbook(EXCEL_PATH, read_only=False, data_only=True)
    ws = wb[SHEET_NAME]
    headers = [cell.value for cell in ws[1]]
    link_idx = headers.index("链接") + 1
    results = []
    for row_index in range(2, ws.max_row + 1):
        cell = ws.cell(row_index, link_idx)
        link = cell.hyperlink.target if cell.hyperlink else cell.value
        if not link:
            continue
        results.append(str(link).strip())
        if limit and len(results) >= limit:
            break
    return results


def topic_id_from_link(link):
    return link.rstrip("/").split("/")[-1]


def fetch_json(tab_id):
    script = "return await fetch(location.pathname + '.json', {credentials: 'include'}).then(r => r.text())"
    raw = run_cmd([AGENT_BROWSER_CLI, "exec", "--tab", str(tab_id), script])
    payload = json.loads(raw)
    result = payload.get("result", {}).get("js_return", "")
    if not isinstance(result, str) or not result.startswith("{"):
        raise RuntimeError(f"Unexpected fetch result: {result[:120]}")
    return result


def open_link(link, topic_id):
    payload = json.loads(run_cmd([AGENT_BROWSER_CLI, "open", link]))
    result = payload.get("result", {})
    tab_id = result.get("opened_tab_id") or result.get("metadata", {}).get("executor", {}).get("tab_id")
    if tab_id:
        return str(tab_id)
    time.sleep(2)
    return str(find_tab_id_by_topic_id(topic_id))


def find_tab_id_by_topic_id(topic_id):
    for _ in range(5):
        tabs_payload = json.loads(run_cmd([AGENT_BROWSER_CLI, "scan", "--tabs-only"]))
        tabs = tabs_payload["result"]["metadata"]["tabs"]
        for tab in reversed(tabs):
            url = str(tab.get("url", ""))
            if f"/{topic_id}" in url:
                return tab.get("tab_id") or tab.get("id")
        time.sleep(1)
    raise RuntimeError(f"Tab not found for topic id: {topic_id}")


def fetch_json_with_retry(tab_id, topic_id):
    try:
        return fetch_json(tab_id)
    except Exception:
        fallback_tab_id = find_tab_id_by_topic_id(topic_id)
        if str(fallback_tab_id) != str(tab_id):
            time.sleep(1)
        return fetch_json(fallback_tab_id)


def main():
    limit = int(sys.argv[1]) if len(sys.argv) > 1 else 10
    CACHE_DIR.mkdir(parents=True, exist_ok=True)
    links = read_links(limit)
    print(f"links={len(links)}")

    for index, link in enumerate(links, start=1):
        topic_id = topic_id_from_link(link)
        cache_file = CACHE_DIR / f"{topic_id}.json"
        if cache_file.exists():
            print(f"[{index}/{len(links)}] skip cache {topic_id}")
            continue

        print(f"[{index}/{len(links)}] open {link}")
        try:
            tab_id = open_link(link, topic_id)
            time.sleep(2)
            json_text = fetch_json_with_retry(tab_id, topic_id)
            cache_file.write_text(json_text, encoding="utf-8")
            print(f"[{index}/{len(links)}] saved {cache_file.name}")
        except Exception as exc:
            print(f"[{index}/{len(links)}] failed {topic_id}: {exc}")


if __name__ == "__main__":
    main()
