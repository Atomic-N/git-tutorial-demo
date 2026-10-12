from pathlib import Path
import shutil

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt


TEMPLATE = Path(
    "D:\\Desktop\\04 A242371174+"
    "\u8f6f\u4ef6\u7cfb\u7edf\u5f00\u53d1\u5b9e\u8bad-"
    "\u8bfe\u7a0b\u8003\u6838\u8bc4\u5206\u7ec6\u5219-"
    "\u653e\u5728\u8bfe\u7a0b\u4f5c\u54c1\u6587\u6863\u7684\u6700\u540e.docx"
)
OUTPUT = TEMPLATE.with_name(
    "04 A242371174+"
    "\u8f6f\u4ef6\u7cfb\u7edf\u5f00\u53d1\u5b9e\u8bad-"
    "\u8bfe\u7a0b\u8003\u6838\u8bc4\u5206\u7ec6\u5219-"
    "\u5df2\u8bc4\u9605-95\u5206.docx"
)

STUDENT_IDS = (
    "202401150631\u3001202401150610\u3001202401150647\u3001"
    "202401150629\u3001202401150602"
)
STUDENT_NAMES = (
    "\u4e8e\u6cfd\u6d9b\u3001\u8c2d\u51cc\u3001\u7f57\u632f\u4e1c\u3001"
    "\u79b9\u5029\u3001\u6b27\u9633\u6b23\u73a5"
)

REVIEW = (
    "\u8be5\u8bfe\u7a0b\u4f5c\u54c1\u4ee5\u201cAI\u6280\u672f\u4ea4\u6d41"
    "\u7efc\u5408\u8bba\u575b\u201d\u4e3a\u9898\uff0c\u9009\u9898\u8d34\u8fd1"
    "\u5f53\u524dAI\u6280\u672f\u5b66\u4e60\u4e0e\u4ea4\u6d41\u9700\u6c42\uff0c"
    "\u9879\u76ee\u5b9a\u4f4d\u660e\u786e\uff0c\u9700\u6c42\u5206\u6790\u80fd"
    "\u591f\u4ece\u7814\u7a76\u80cc\u666f\u3001\u7528\u6237\u89d2\u8272\u3001"
    "\u524d\u540e\u53f0\u529f\u80fd\u3001\u4e1a\u52a1\u6d41\u7a0b\u53ca\u975e"
    "\u529f\u80fd\u9700\u6c42\u7b49\u65b9\u9762\u5c55\u5f00\uff0c\u5185\u5bb9"
    "\u8f83\u5168\u9762\uff0c\u4e0e\u6700\u7ec8\u7cfb\u7edf\u529f\u80fd\u5bf9"
    "\u5e94\u5173\u7cfb\u6e05\u6670\u3002\u7cfb\u7edf\u8bbe\u8ba1\u9075\u5faa"
    "B/S\u67b6\u6784\u3001MVC\u8bbe\u8ba1\u6a21\u5f0f\u548c\u4e09\u5c42"
    "\u7ed3\u6784\u601d\u60f3\uff0c\u6a21\u5757\u5212\u5206\u5408\u7406\uff0c"
    "\u6570\u636e\u5e93\u6982\u5ff5\u8bbe\u8ba1\u4e0e\u7269\u7406\u7ed3\u6784"
    "\u8bbe\u8ba1\u8f83\u5b8c\u6574\uff0c\u80fd\u591f\u652f\u6491\u8bba\u575b"
    "\u7684\u7528\u6237\u3001\u5206\u7c7b\u6807\u7b7e\u3001\u8bdd\u9898\u56de"
    "\u590d\u3001\u901a\u77e5\u79c1\u4fe1\u3001\u5185\u5bb9\u6cbb\u7406\u548c"
    "\u540e\u53f0\u7ba1\u7406\u7b49\u6838\u5fc3\u4e1a\u52a1\u3002\u7cfb\u7edf"
    "\u5b9e\u73b0\u90e8\u5206\u4f53\u73b0\u4e86\u8f83\u597d\u7684\u5de5\u7a0b"
    "\u5b9e\u8df5\u80fd\u529b\uff0c\u91c7\u7528JSP\u3001Servlet\u3001"
    "JavaBean\u3001JDBC\u3001MySQL\u7b49\u6280\u672f\u5b8c\u6210\u4e3b\u8981"
    "\u529f\u80fd\uff0c\u524d\u540e\u53f0\u4e1a\u52a1\u95ed\u73af\u8f83\u5b8c"
    "\u6574\uff0c\u4ee3\u7801\u7ec4\u7ec7\u4e0e\u529f\u80fd\u8bf4\u660e\u8f83"
    "\u89c4\u8303\u3002\u6d4b\u8bd5\u4e0e\u603b\u7ed3\u90e8\u5206\u5305\u542b"
    "\u6d4b\u8bd5\u73af\u5883\u3001\u6027\u80fd\u6d4b\u8bd5\u3001\u6ce8\u518c"
    "\u767b\u5f55\u3001\u5206\u7c7b\u6d4f\u89c8\u3001\u8bdd\u9898\u53d1\u5e03"
    "\u4e0e\u56de\u590d\u3001\u4e92\u52a8\u529f\u80fd\u3001\u540e\u53f0\u7ba1"
    "\u7406\u7b49\u6d4b\u8bd5\u5185\u5bb9\uff0c\u80fd\u591f\u8f83\u597d\u9a8c"
    "\u8bc1\u7cfb\u7edf\u53ef\u7528\u6027\uff0c\u5e76\u5bf9\u9879\u76ee\u5f00"
    "\u53d1\u8fc7\u7a0b\u8fdb\u884c\u4e86\u603b\u7ed3\u53cd\u601d\u3002\u603b"
    "\u4f53\u6765\u770b\uff0c\u4f5c\u54c1\u6587\u6863\u7ed3\u6784\u5b8c\u6574"
    "\u3001\u5185\u5bb9\u5145\u5b9e\u3001\u6280\u672f\u8def\u7ebf\u6e05\u695a"
    "\uff0c\u7cfb\u7edf\u529f\u80fd\u5b8c\u6210\u5ea6\u9ad8\uff0c\u8fbe\u5230"
    "\u8bfe\u7a0b\u5b9e\u8bad\u4f18\u79c0\u6c34\u5e73\u3002\u540e\u7eed\u53ef"
    "\u8fdb\u4e00\u6b65\u8865\u5145\u5f02\u5e38\u573a\u666f\u3001\u538b\u529b"
    "\u6d4b\u8bd5\u6570\u636e\u548c\u90e8\u7f72\u8fd0\u7ef4\u7ec6\u8282\uff0c"
    "\u4f7f\u6d4b\u8bd5\u8bc1\u636e\u4e0e\u5de5\u7a0b\u8bf4\u660e\u66f4\u52a0"
    "\u5145\u5206\u3002\u7efc\u5408\u8bc4\u5b9a\u4e3a\u4f18\u79c0\uff0c\u6210"
    "\u7ee995\u5206\u3002"
)


def set_font(run, size=10.5, bold=False):
    run.font.name = "\u5b8b\u4f53"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "\u5b8b\u4f53")
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold


def clear_cell(cell):
    for para in cell.paragraphs:
        para.clear()
    for para in cell.paragraphs[1:]:
        para._element.getparent().remove(para._element)
    return cell.paragraphs[0]


def set_cell_text(cell, text, size=10.5, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER):
    para = clear_cell(cell)
    para.alignment = align
    para.paragraph_format.space_before = Pt(0)
    para.paragraph_format.space_after = Pt(0)
    run = para.add_run(text)
    set_font(run, size=size, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def replace_paragraph(paragraph, text):
    paragraph.clear()
    run = paragraph.add_run(text)
    set_font(run, size=10.5)


def main():
    shutil.copyfile(TEMPLATE, OUTPUT)
    doc = Document(str(OUTPUT))

    for paragraph in doc.paragraphs:
        if paragraph.text.startswith("\u8003\u751f\u5b66\u53f7\uff1a"):
            replace_paragraph(paragraph, f"\u8003\u751f\u5b66\u53f7\uff1a{STUDENT_IDS}")
        elif paragraph.text.startswith("\u8003\u751f\u59d3\u540d\uff1a"):
            replace_paragraph(paragraph, f"\u8003\u751f\u59d3\u540d\uff1a {STUDENT_NAMES}")

    for idx, score in enumerate(["19", "19", "38", "19", "95"], start=1):
        set_cell_text(doc.tables[0].cell(1, idx), score, size=11, bold=True)

    review_cell = doc.tables[1].cell(4, 1)
    paragraph = clear_cell(review_cell)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.first_line_indent = Pt(21)
    paragraph.paragraph_format.line_spacing = 1.2
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(REVIEW)
    set_font(run, size=10.5)

    sign = review_cell.add_paragraph()
    sign.alignment = WD_ALIGN_PARAGRAPH.LEFT
    sign.paragraph_format.space_before = Pt(8)
    sign.paragraph_format.space_after = Pt(0)
    run = sign.add_run("\u6307\u5bfc\u8001\u5e08\u7b7e\u540d\uff1a")
    set_font(run, size=10.5)

    date = review_cell.add_paragraph()
    date.alignment = WD_ALIGN_PARAGRAPH.LEFT
    date.paragraph_format.space_before = Pt(10)
    date.paragraph_format.space_after = Pt(0)
    run = date.add_run("\u5e74   \u6708   \u65e5")
    set_font(run, size=10.5)
    review_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    doc.save(str(OUTPUT))
    print(str(OUTPUT))


if __name__ == "__main__":
    main()
